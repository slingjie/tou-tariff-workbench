"""Cross-check our official-source tariff data (data/tariff.db, READ-ONLY) against energydc.cn.

energydc.cn is CROSS-CHECK ONLY. Its ToS §4 forbids reuse ("未经授权…不得擅自使用、复制、修改或传播"), so:
  * nothing from energydc is ever written to data/tariff.db or data/json/;
  * comparison outputs go to work/compare/ for private review only;
  * raw API responses are cached under data/raw/energydc/ (policy_<uuid>.json) so reruns make no new requests.

Sources (guest-public only, same XHR the public /policy page makes; no login, no paid/gated endpoints):
  sitemap  https://energydc.cn/sitemap/sitemap_policy_1.xml
  detail   https://energydc.cn/edc/api/v1/eprice/policy/<uuid>
Fetching is delegated to pipeline/sources/energydc_api_probe.py (>=2 s between requests, <=2 retries).

Usage:
  python -m pipeline.compare_energydc --region jiangsu [--region shanghai ...]   # rerunnable per region
  python -m pipeline.compare_energydc --all-loaded                               # every region in tariff.db
  python -m pipeline.compare_energydc --rebuild-summary                          # only regenerate summary/xlsx
Options: --months 2025-11..2026-10 (default), --no-fetch (use cache only), --budget SECONDS (stop fetching early)

Outputs (work/compare/): <region>.csv (all diffs/unmatched/coverage/TOU lines), <region>.json (summary row),
summary.csv, compare_report.xlsx (汇总, 差异明细, TOU). Progress appended to work/progress_compare.md.
"""
import argparse, csv, glob, json, os, re, sqlite3, sys, time
from datetime import datetime, timezone, timedelta

from . import parse_common as pc
from .sources import energydc_api_probe as ep

ROOT = pc.ROOT
OUT = os.path.join(ROOT, "work", "compare")
PROGRESS = os.path.join(ROOT, "work", "progress_compare.md")
CST = timezone(timedelta(hours=8))

# our region code -> (energydc policy slug, caption regex that selects the sub-region table or None)
REGION_MAP = {
    "jiangsu": ("jiangsu", None), "shanghai": ("shanghai", None), "fujian": ("fujian", None),
    "hebei-south": ("hebei-south", None), "jibei": ("hebei-jibei", None), "ningxia": ("ningxia", None),
    "tianjin": ("tianjin", None), "liaoning": ("liaoning", None), "jiangxi": ("jiangxi", None),
    "chongqing": ("chongqing", None), "sichuan": ("sichuan", None), "zhejiang": ("zhejiang", None),
    "anhui": ("anhui", None), "mengxi": ("mengxi-grid", None), "mengdong": ("mengdong-grid", None),
    # from 2026-08 Guangdong publishes 珠三角6市 (= 珠三角五市 + 江门) – patterns are tried in order
    "guangdong-prd5": ("guangdong", [r"珠三角(五|5)市", r"珠三角(6|六)市"]), "guangdong-huizhou": ("guangdong", r"惠州"),
    "guangdong-jiangmen": ("guangdong", [r"江门", r"珠三角(6|六)市"]), "guangdong-eastwest": ("guangdong", r"东西两翼"),
    "guangdong-north": ("guangdong", r"粤北"),
    "shenzhen": ("shenzhen", r"^(?!.*深汕)"), "shenzhen-shenshan": ("shenzhen", r"深汕"),
    "shanxi": ("shanxi", None), "beijing": ("beijing", None), "heilongjiang": ("heilongjiang", None),
    "jilin": ("jilin", None), "shandong": ("shandong", None), "henan": ("henan", None),
    "hubei": ("hubei", None), "hunan": ("hunan", None),
    # energydc captions: '…电价表（不含榆林地区）' vs '…电价表（榆林地区）' – the old r"榆林" matched both (and the
    # negative look-ahead excluded 不含榆林), so yulin was compared against both tables. From 2026-08 the 不含榆林 table
    # has no caption at all → the "only unnamed table" fallback in edc_rows_for picks it.
    "shaanxi": ("shaanxi", r"不含榆林"), "shaanxi-yulin": ("shaanxi", r"(?<!不含)榆林"),
    "gansu": ("gansu", None), "qinghai": ("qinghai", None), "xinjiang": ("xinjiang", None),
    "guangxi": ("guangxi", None), "yunnan": ("yunnan", None), "guizhou": ("guizhou", None),
    "hainan": ("hainan", None),
}
EXCLUDE_TABLE = re.compile(r"1\.5|1．5|倍|转供|充换电|充电|电动汽车")
FIELDS = ["energy_price", "purchase_price", "line_loss", "tdp", "system_op", "gov_fund",
          "sharp", "peak", "flat", "valley", "deep_valley", "demand_price", "capacity_price"]
FIELD_CN = {"energy_price": "电度电价", "purchase_price": "代理购电价格", "line_loss": "线损", "tdp": "输配电价",
            "system_op": "系统运行费", "gov_fund": "政府性基金", "sharp": "尖峰", "peak": "高峰", "flat": "平段",
            "valley": "低谷", "deep_valley": "深谷", "demand_price": "需量电价", "capacity_price": "容量电价"}
TOL = {f: 0.00005 for f in FIELDS}
TOL["demand_price"] = TOL["capacity_price"] = 0.05


def month_range(spec):
    a, b = spec.split("..")
    y, m = map(int, a.split("-"))
    out = []
    while f"{y:04d}-{m:02d}" <= b:
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def log(msg):
    os.makedirs(os.path.dirname(PROGRESS), exist_ok=True)
    new = not os.path.exists(PROGRESS)
    with open(PROGRESS, "a", encoding="utf-8") as f:
        if new:
            f.write("# energydc cross-check progress (pipeline/compare_energydc.py)\n\n")
        f.write(f"- {datetime.now(CST).strftime('%Y-%m-%d %H:%M')} {msg}\n")
    print(msg)


# ---------------------------------------------------------------- energydc side
def edc_pages(slug, month, fetch=True, deadline=None):
    """-> list of policy dicts (typeId 2 only) for slug/month, with '_uuid'. Uses cache; fetches if allowed."""
    out = []
    for u in ep.find_uuids(slug, month):
        path = os.path.join(ep.CACHE, f"policy_{u}.json")
        if not os.path.exists(path):
            if not fetch or (deadline and time.time() > deadline):
                out.append({"_uuid": u, "_missing": True})
                continue
            try:
                ep.fetch_policy(u)
            except Exception as e:  # probe already did <=2 retries
                out.append({"_uuid": u, "_missing": True, "_error": str(e)[:200]})
                continue
        d = json.load(open(path, encoding="utf-8"))
        d["_uuid"] = u
        if d.get("typeId") == 2:
            out.append(d)
    return out


def title_rank(t):
    """lower = preferred. User rule: prefer the 代理购电工商业用户电价表 page by title."""
    t = t or ""
    r = 0
    if not re.search(r"代理购电工商业用户电价表", t):
        r += 1
    if re.search(r"周|修正", t):
        r += 4
    if re.search(r"实际结算", t):
        r += 2
    return r


def _fix_md(md):
    # energydc text truncation artefacts inside cells: "1-10（20）千" / "220千伏及" / "220千伏及以"
    md = re.sub(r"(\d[)）]?)千\s*\|", r"\1千伏 |", md)
    md = re.sub(r"千伏及(以)?\s*\|", "千伏及以上 |", md)
    return md


VOLT_CELL = re.compile(r"千伏(?!安)|kV(?!A)|伏$|电压等级")
HDR_FIRST = ("用电分类", "用电类别", "计算关系", "公式", "")


def _is_header(row):
    h = pc.norm(row[0])
    return h in HDR_FIRST or set("".join(row)) <= set("-: ")


def parse_table(rows):
    """Header-driven parser (adapted from energydc_api_probe.parse_table) with fixes:
    '用电类别' header (Shenzhen), tables without an explicit 电压等级 column, split purchase_price columns
    (Ningxia: 高峰/平/低谷 → take 平; Shandong sub-components → sum), extra qualifier label (kVA / kWh band)."""
    hdr = [r for r in rows if _is_header(r)]
    data = [r for r in rows if not _is_header(r)]
    if not hdr or not data:
        return []
    ncol = max(len(r) for r in rows)
    merged = []
    for j in range(ncol):
        parts = []
        for r in hdr:
            if pc.norm(r[0]) in ("计算关系", "公式"):
                continue
            c = r[j] if j < len(r) else ""
            if c and c != "nan" and set(c) - set("-: ") and c not in parts:
                parts.append(c)
        merged.append(pc.norm("".join(parts)))
    alltext = "".join(merged)
    if "谷" not in alltext:
        return []
    in_fen = "分/千瓦时" in alltext
    vcol = next((j for j, h in enumerate(merged) if "电压等级" in h), None)
    if vcol is None:  # label columns = everything before the first price column
        first_price = next((j for j, h in enumerate(merged) if re.search(r"电价|电量|电度|价格", h)), None)
        if first_price is None:
            return []
        vcol = first_price - 1
        # Shenzhen: the column right before the price block can be a nameless label column
    fieldcols = {}
    for j, h in enumerate(merged):
        if j <= vcol:
            continue
        hh = re.sub(r"其中|(?<!非)分时电[度量]用?电价(格)?|容\(需\)量用?电价(格)?|\(.*?\)|（.*?）", "", h) or h
        if re.search(r"代购电价格|代理购电价格", h) and re.search(r"峰|平|谷", hh) and not re.search(r"分时电度", h):
            fieldcols.setdefault("purchase_price", []).append((j, h))
            continue
        for f, inc, exc in ep.COLMAP:
            if re.search(inc, hh) and not (exc and re.search(exc, hh)):
                fieldcols.setdefault(f, []).append((j, h))
                break
    # data label columns may extend beyond vcol when header has 'nan' label columns (Shenzhen band column)
    out = []
    for r in data:
        r = [("" if c == "nan" else c) for c in r] + [""] * (ncol - len(r))
        # label cells: up to vcol, plus any non-numeric cells right after it that sit in an unnamed column
        label = [c for c in r[:vcol + 1] if c]
        j = vcol + 1
        while j < ncol and r[j] and pc.num(r[j].replace("\\", "")) is None and r[j] not in ("/", "-", "—", "／"):
            label.append(r[j]); j += 1
        extra = j - (vcol + 1)
        if extra and merged[vcol + 1]:
            # extra label column(s) with no header of their own (Shenzhen 2025-11: kWh band) → values shifted right
            r = r[:vcol + 1] + r[vcol + 1 + extra:] + [""] * extra
        volt = next((c for c in label if VOLT_CELL.search(pc.norm(c)) and not re.search(r"kW·h|千瓦时", c)), None)
        if not volt:
            continue
        mm = re.match(r"^(.*?用电\s*[(（][^)）]*[)）])\s*(.*?)(?:\s*[-－]\s*(\d+\s*kW.*))?$", volt)
        if mm and mm.group(2):  # energydc packed '一、工商业用电 (101至3000kVA) 10千伏高供高计 - 250kWh…' in one cell
            label = [mm.group(1), mm.group(2)] + ([mm.group(3)] if mm.group(3) else []) + [c for c in label if c != volt]
            volt = mm.group(2)
        joined = "".join(label)
        tariff = "两部制" if "两部制" in joined else ("单一制" if "单一制" in joined else None)
        cat = next((c for c in label if c != volt and ("用电" in c or "工业" in c or "工商业" in c or "用户" in c)), None)
        qual = [c for c in label if c not in (volt, cat) and re.search(r"千伏安|kW·h|kVA", c)]
        cat = re.sub(r"[-－]?\s*(单一制|两部制)", "", cat or "").strip() or None
        row = {"category": cat, "tariff_type": tariff, "voltage": volt, "qual": " ".join(qual)}
        for f, cols in fieldcols.items():
            vals = []
            if f == "purchase_price" and len(cols) > 1:
                flat_cols = [c for c, h in cols if re.search(r"平", h)]
                use = flat_cols[:1] if flat_cols else [c for c, _ in cols]
            else:
                use = [c for c, _ in cols]
            for c in use:
                v = pc.num(r[c].replace("\\", ""))
                if v is not None:
                    vals.append(v)
            v = (sum(vals) if (f in ep.SUM_FIELDS and len(use) > 1) else vals[0]) if vals else None
            if v is not None and in_fen and f in ep.ENERGY_FIELDS:
                v = v / 100.0
            row[f] = round(v, 8) if v is not None else None
        if row["tariff_type"] is None:  # Shenzhen-style: no 单一制/两部制 label → infer from demand/capacity
            row["tariff_type"] = "两部制" if (row.get("demand_price") or row.get("capacity_price")) else "单一制"
            row["tariff_inferred"] = True
        out.append(row)
    # merged band cell (Shenzhen 2025-11): second row of a 及以下/以上 pair has an empty band → complement
    for a, b in zip(out, out[1:]):
        if not b["qual"] and a["qual"] and a["category"] == b["category"] and a["voltage"] == b["voltage"]:
            b["qual"] = a["qual"].replace("及以下", "以上") if "及以下" in a["qual"] else ""
            b["qual_inferred"] = True
    # unit guard: header said nothing usable about 分/千瓦时 but values are clearly in 分 (e.g. Shenzhen 2026-10)
    probe = [r.get("energy_price") or r.get("flat") for r in out if (r.get("energy_price") or r.get("flat"))]
    if probe and not in_fen and sorted(probe)[len(probe) // 2] > 5:
        for r in out:
            for f in ep.ENERGY_FIELDS:
                if r.get(f) is not None:
                    r[f] = round(r[f] / 100.0, 8)
            r["unit_guard"] = True
    return out


SUBREGION_NAMES = r"珠三角|惠州|江门|东西两翼|粤北|深汕|深圳市|榆林"


def md_tables(md):
    """like energydc_api_probe._tables, but the caption also carries the nearest markdown heading above the
    table (energydc sometimes puts '## …（惠州市）' several lines above the table)."""
    out, cur, prev, cap, heading = [], [], [], "", ""
    for ln in md.splitlines():
        if ln.lstrip().startswith("|"):
            if not cur:
                cap = (heading + " / " if heading else "") + " / ".join(prev[-2:])
            cur.append(ep._cells(ln))
        else:
            if cur:
                out.append((cap, cur)); cur = []; heading = ""
            s = ln.strip()
            if s.startswith("#"):
                heading = s[:120]
            if s:
                prev.append(s[:120])
    if cur:
        out.append((cap, cur))
    return out


def edc_rows_for(page, caption_re, slug=None):
    """-> (rows, captions, note). Excludes 1.5x / 转供电 / 充换电 tables; selects the sub-region table by caption
    (list of patterns tried in order); if no caption matches and exactly one table carries no sub-region name
    at all, that one is used (energydc dropped its heading)."""
    tables = []
    for cap, t in md_tables(_fix_md(page.get("content", ""))):
        if EXCLUDE_TABLE.search(cap):
            continue
        parsed = [r for r in parse_table(t) if not EXCLUDE_TABLE.search(r.get("category") or "")]
        if parsed:
            tables.append((cap, parsed))
    note = ""
    if caption_re:
        pats = caption_re if isinstance(caption_re, list) else [caption_re]
        sel = []
        for pat in pats:
            sel = [(c, p) for c, p in tables if re.search(pat, c)]
            if sel:
                break
        if not sel and tables and any(re.search(SUBREGION_NAMES, c) for c, _ in tables):
            unnamed = [(c, p) for c, p in tables if not re.search(SUBREGION_NAMES, c)]
            if len(unnamed) == 1:
                sel = unnamed
                note = "sub-region table = the only table without a sub-region heading (energydc dropped it)"
        tables = sel
    rows, caps = [], []
    for cap, parsed in tables:
        caps.append(cap)
        for r in parsed:
            r["table_caption"] = cap
        rows += parsed
    return rows, caps, note


# ---------------------------------------------------------------- normalisation / matching
def volt_canon(v):
    s = pc.norm(v).replace("－", "-").replace("~", "-").replace("—", "-").replace("–", "-").replace("至", "-")
    s = re.sub(r"\((20|66)\)", "", s)
    s = re.sub(r"\(不含\)", "", s)
    s = s.replace("千伏", "kV").replace("KV", "kV")
    s = re.sub(r"kV-(\d)", r"-\1", s)            # 35kV-110kV -> 35-110kV
    s = s.replace("及以上", "+").replace("以上", "+")
    s = re.sub(r"[及以]$", "+", s)
    return s


def qual_canon(*texts):
    t = pc.norm(" ".join(x or "" for x in texts)).replace("／", "/").replace("千瓦安", "千伏安")  # hainan edc typo 2026-06
    m = re.search(r"(\d+)kW·?h(及以下|及以上|以上|以下)", t)
    if m:
        return f"{m.group(1)}kWh" + {"及以上": "以上", "以下": "及以下"}.get(m.group(2), m.group(2))
    m = re.search(r"100千伏安(及以上|以下|及以下)", t)
    if m:
        return f"100kVA{m.group(1)}"
    return ""


def cat_class(c):
    c = pc.norm(c)
    m = re.match(r"^([一二三四五])、", c)
    if m:
        return m.group(1)
    if "大工业" in c and "一般" not in c:
        return "大工业"
    if "一般" in c and "大工业" not in c:
        return "一般"
    return ""


def keyset(r, ours):
    rank = pc.voltage_rank(r["voltage"])
    q = qual_canon(r.get("category"), r.get("voltage"), r.get("qual"))
    cc = cat_class(r.get("category"))
    v = volt_canon(re.sub(r"\[.*?\]", "", r["voltage"] or ""))
    t = r.get("tariff_type")
    return [("L1", (t, cc, q, v)), ("L2", (t, cc, q, rank)), ("L3", (t, q, v)), ("L4", (t, q, rank))]


def match_rows(ours, edc):
    """layered fuzzy matching; a key is used only if it is unique on both sides among still-unmatched rows."""
    pairs, mo, me = [], set(), set()
    for lvl in range(4):
        idx_o, idx_e = {}, {}
        for i, r in enumerate(ours):
            if i not in mo:
                idx_o.setdefault(keyset(r, True)[lvl][1], []).append(i)
        for i, r in enumerate(edc):
            if i not in me:
                idx_e.setdefault(keyset(r, False)[lvl][1], []).append(i)
        for k, lo in idx_o.items():
            le = idx_e.get(k)
            if le and len(lo) == 1 and len(le) == 1 and None not in (k[-1],):
                pairs.append((lo[0], le[0], f"L{lvl+1}"))
                mo.add(lo[0]); me.add(le[0])
    # L5: the only remaining row on each side in the same category class (e.g. Shenzhen 第三类 不分电压等级)
    ro = [i for i in range(len(ours)) if i not in mo]
    re_ = [i for i in range(len(edc)) if i not in me]
    for cc in set(cat_class(ours[i]["category"]) for i in ro):
        lo = [i for i in ro if cat_class(ours[i]["category"]) == cc]
        le = [i for i in re_ if cat_class(edc[i]["category"]) == cc]
        if cc and len(lo) == 1 and len(le) == 1:
            pairs.append((lo[0], le[0], "L5")); mo.add(lo[0]); me.add(le[0])
    return pairs, [i for i in range(len(ours)) if i not in mo], [i for i in range(len(edc)) if i not in me]


def compare_pair(o, e):
    out = []
    for f in FIELDS:
        a, b = o.get(f), e.get(f)
        if a is None and b is None:
            continue
        if a is None or b is None:
            kind = "ours_null" if a is None else "edc_null"
            # known energydc artefact: merged-cell propagation of 需量/容量 into 单一制 rows
            if a is None and f in ("demand_price", "capacity_price") and o.get("tariff_type") == "单一制":
                kind = "edc_merged_cell_artefact"
            out.append((f, a, b, None, kind))
        elif abs(a - b) > TOL[f] + 1e-12:
            out.append((f, a, b, round(abs(a - b), 8), "value"))
    return out


# ---------------------------------------------------------------- TOU
DOCNO = re.compile(r"[\u4e00-\u9fa5]{1,10}〔\s*\d{4}\s*〕\s*\d+\s*号")
_DOC_PREFIX = re.compile(r"^.*?(?:根据|依据|按照|参照|执行|据|照|按|依|与|和|及|、)(?=[\u4e00-\u9fa5]{2,8}〔)")


def cited_tou_docs(content):
    """文号 cited within ~60 chars of a 分时/峰谷 mention (best effort; may include generic pricing docs)."""
    text = re.sub(r"\s+", "", content or "")
    docs = []
    for m in DOCNO.finditer(text):
        win = text[max(0, m.start() - 60): m.end() + 30]
        if not re.search(r"分时|峰谷|尖峰|深谷", win):
            continue
        d = _DOC_PREFIX.sub("", m.group(0))
        if d not in docs:
            docs.append(d)
    return docs


def _tmin(t):
    h, m = t.split(":")
    return int(h) * 60 + int(m)


def _segs_norm(segs):
    """split wrap-around segments at midnight, merge adjacent same-kind segments, normalise 24:00."""
    raw = []
    for a, b, k in segs:
        a = a.zfill(5); b = b.zfill(5)
        if b in ("00:00", "24:00"):
            b = "24:00"
        if _tmin(b) <= _tmin(a) and b != "24:00":
            raw += [[a, "24:00", k], ["00:00", b, k]]
        else:
            raw.append([a, b, k])
    raw.sort(key=lambda x: _tmin(x[0]))
    out = []
    for a, b, k in raw:
        if out and out[-1][2] == k and out[-1][1] == a:
            out[-1][1] = b
        else:
            out.append([a, b, k])
    return out


_KIND = {"峰": "peak", "谷": "valley", "尖峰": "sharp", "高峰": "peak", "平段": "flat", "平时段": "flat", "平时": "flat", "谷段": "valley", "平": "flat", "低谷": "valley", "深谷": "deep_valley"}
_T = r"(?:次日)?\d{1,2}(?::\d{2})?(?:点|时)?"
_SEG = re.compile(r"(尖峰|高峰|平时段|平时|平段|平|低谷|深谷|谷段|(?<![高低尖深])峰(?=时段)|(?<![高低尖深])谷(?=时段))(?:时段)?(?:为|:)?((?:" + _T + r"-" + _T + r"[、,，和]?)+)")
_MITEM = r"(?:次年)?\d{1,2}月?(?:-(?:次年)?\d{1,2}月?)?"
_HDR = re.compile(r"(?<![\d.〕:])((?:" + _MITEM + r")(?:[、，,和及]" + _MITEM + r")*(?:份)?|其他月份|其余月份|其他季节|其他时间)"
                  r"(?=[)）]?[^。；;:]{0,24}?(?:时段划分|时段设置|时段为|时段:|分时电价时段|每日|尖峰|高峰|平段|平时|低谷|深谷|谷段|峰时段|谷时段))")


def _hm(x):
    x = x.replace("次日", "").replace("点", "").replace("时", "")
    h, _, m = x.partition(":")
    return f"{int(h):02d}:{int(m or 0):02d}"


def _months_list(h):
    if re.match(r"其[他余]", h):
        return ["REST"]
    ms = []
    for a, b in re.findall(r"(?:次年)?(\d{1,2})月?(?:-(?:次年)?(\d{1,2})月?)?", h):
        a = int(a); b = int(b) if b else a
        ms += list(range(a, b + 1)) if a <= b else list(range(a, 13)) + list(range(1, b + 1))
    return sorted(set(m for m in ms if 1 <= m <= 12))


def parse_tou_text(content):
    """Best-effort TOU periods from the announcement notes (free text). Returns {month: segments} using the
    first clause that covers each month (the main table's note comes first; EV-charging notes come later).
    Handles '高峰时段：8:00-11:00、13:00-17:00', '平段时段：', '高峰时段为10-12点', '其余时段为平段'."""
    t = re.sub(r"\s+", "", content or "")
    t = t.replace("：", ":").replace("—", "-").replace("－", "-").replace("~", "-").replace("～", "-").replace("至", "-")
    out = {}
    for sent in re.split(r"。", t):
        if not _SEG.search(sent) or not re.search(r"峰", sent):
            continue
        if re.search(r"充换电|充电设施|电动汽车|储能", sent[:60]):
            continue
        hdrs = [(m.start(), m.end(), _months_list(m.group(1))) for m in _HDR.finditer(sent)
                if "月" in m.group(1) or re.match(r"其[他余]", m.group(1))]
        hdrs = [h for h in hdrs if h[2]]
        spans = []
        if hdrs:
            for i, (a, b, ms) in enumerate(hdrs):
                end = hdrs[i + 1][0] if i + 1 < len(hdrs) else len(sent)
                spans.append((ms, sent[b:end]))
        else:
            spans.append((list(range(1, 13)), sent))
        for ms, body in spans:
            segs = []
            for k, times in _SEG.findall(body):
                for a, b in re.findall(r"(" + _T + r")-(" + _T + r")", times):
                    segs.append([_hm(a), _hm(b), _KIND[k]])
            if not segs or not any(k in ("peak", "sharp") for _, _, k in segs):
                continue
            if re.search(r"其余时段为平|其他时段为平|其余为平", body) or not any(k == "flat" for _, _, k in segs):
                segs = _fill_flat(segs)
            if ms == ["REST"]:
                ms = [m for m in range(1, 13) if m not in out]
            for m in ms:
                out.setdefault(m, segs)
    return out


def _fill_flat(segs):
    cov = [None] * 1440
    for a, b, k in _segs_norm(segs):
        for i in range(_tmin(a), _tmin(b)):
            cov[i] = k
    out, i = [], 0
    while i < 1440:
        j = i
        while j < 1440 and cov[j] is None:
            j += 1
        if j > i:
            out.append([f"{i//60:02d}:{i%60:02d}", f"{j//60:02d}:{j%60:02d}" if j < 1440 else "24:00", "flat"])
        i = j + 1 if j == i else j
    return segs + out


def our_tou_for_month(con, region, month):
    first = month + "-01"
    rules = con.execute("SELECT months, day_type, segments, effective_from, effective_to, doc_id FROM tou_rule "
                        "WHERE region_code=? AND (effective_from IS NULL OR effective_from<=?) "
                        "AND (effective_to IS NULL OR effective_to>=?)", (region, first, first)).fetchall()
    m = int(month[5:])
    hit = [r for r in rules if m in json.loads(r[0]) and r[1] == "normal"]
    if not hit:
        return None, None
    hit.sort(key=lambda r: r[3] or "", reverse=True)  # latest effective wins
    return _segs_norm(json.loads(hit[0][2])), hit[0][5]


# ---------------------------------------------------------------- main per-region
def run_region(region, months, fetch=True, budget=None):
    if region not in REGION_MAP:
        log(f"**{region}**: SKIP – no energydc slug mapping")
        return None
    slug, cap_re = REGION_MAP[region]
    deadline = time.time() + budget if budget else None
    con = sqlite3.connect(f"file:{pc.DB_PATH}?mode=ro", uri=True, timeout=120)
    con.row_factory = sqlite3.Row
    lines = []  # csv rows
    st = dict(region=region, edc_slug=slug, months_both=0, months_ours_only=[], months_edc_only=[],
              months_neither=[], ours_rows=0, edc_rows=0, matched_rows=0, rows_with_diff=0, value_diffs=0,
              null_mismatch_cells=0, edc_artefact_cells=0, unmatched_ours=0, unmatched_edc=0, max_abs_diff=0.0,
              tou_months_compared=0, tou_months_diff=0, notes=[])
    tou_docs_ours = [f"{r[0]} {r[1] or ''}".strip() for r in con.execute(
        "SELECT DISTINCT p.doc_no, p.title FROM tou_rule t JOIN policy_doc p ON p.id=t.doc_id WHERE t.region_code=?",
        (region,))]
    edc_docs_seen = {}
    base = dict(region=region, month="", check="", category_ours="", tariff_type="", voltage_ours="",
                category_edc="", voltage_edc="", match_level="", field="", field_cn="", ours="", edc="",
                abs_diff="", tol="", diff_type="", edc_title="", edc_url="", note="")
    for month in months:
        ours = [dict(r) for r in con.execute("SELECT * FROM tariff_month WHERE region_code=? AND month=?", (region, month))]
        pages = edc_pages(slug, month, fetch=fetch, deadline=deadline)
        missing = [p for p in pages if p.get("_missing")]
        pages = [p for p in pages if not p.get("_missing")]
        cands = []
        for p in pages:
            rows, caps, tnote = edc_rows_for(p, cap_re, slug)
            if tnote:
                st["notes"].append(f"{month}: {tnote}")
            if rows:
                cands.append((title_rank(p.get("title")), p, rows))
        cands.sort(key=lambda x: x[0])
        if missing and not cands:
            st["notes"].append(f"{month}: {len(missing)} energydc page(s) not fetched")
        has_o, has_e = bool(ours), bool(cands)
        if has_o and has_e:
            st["months_both"] += 1
        elif has_o:
            st["months_ours_only"].append(month)
        elif has_e:
            st["months_edc_only"].append(month)
        else:
            st["months_neither"].append(month)
        cov = "both" if has_o and has_e else ("ours_only" if has_o else ("edc_only" if has_e else "neither"))
        if cov != "both":
            lines.append(dict(base, month=month, check="coverage", diff_type=cov,
                              edc_title=(cands[0][1].get("title") if cands else ""),
                              edc_url=(f"{ep.BASE}/policy/{slug}/{month}/{cands[0][1]['_uuid']}" if cands else ""),
                              note=("energydc pages exist but none parsed for this sub-region" if pages and not cands else "")))
        if not cands:
            continue
        _, page, erows = cands[0]
        # dedupe energydc rows on their L1 key (same table repeated)
        seen, ded = set(), []
        for r in erows:
            k = keyset(r, False)[0][1]
            if k not in seen:
                seen.add(k); ded.append(r)
        erows = ded
        url = f"{ep.BASE}/policy/{slug}/{month}/{page['_uuid']}"
        title = page.get("title")
        if len(cands) > 1:
            alt = []
            for rk, p2, r2 in cands[1:]:
                pr, _, _ = match_rows(ours, r2)
                nd = sum(1 for io, ie, _ in pr for d in compare_pair(ours[io], r2[ie]) if d[4] == "value")
                alt.append(f"'{p2.get('title')}' ({nd} value diffs)")
            pr, _, _ = match_rows(ours, erows)
            nd0 = sum(1 for io, ie, _ in pr for d in compare_pair(ours[io], erows[ie]) if d[4] == "value")
            st["notes"].append(f"{month}: {len(cands)} energydc pages; used '{title}' ({nd0} value diffs); others: " + "; ".join(alt))
        if has_o:
            st["ours_rows"] += len(ours); st["edc_rows"] += len(erows)
            pairs, uo, ue = match_rows(ours, erows)
            st["matched_rows"] += len(pairs)
            for io, ie, lvl in pairs:
                o, e = ours[io], erows[ie]
                diffs = compare_pair(o, e)
                if any(d[4] == "value" for d in diffs):
                    st["rows_with_diff"] += 1
                for f, a, b, ad, kind in diffs:
                    if kind == "value":
                        st["value_diffs"] += 1
                        st["max_abs_diff"] = max(st["max_abs_diff"], ad) if f not in ("demand_price", "capacity_price") else st["max_abs_diff"]
                    elif kind == "edc_merged_cell_artefact":
                        st["edc_artefact_cells"] += 1
                    else:
                        st["null_mismatch_cells"] += 1
                    lines.append(dict(base, month=month, check="price", category_ours=o["category"],
                                      tariff_type=o["tariff_type"], voltage_ours=o["voltage"],
                                      category_edc=(e.get("category") or "") + (f" [{e['qual']}]" if e.get("qual") else ""),
                                      voltage_edc=e["voltage"], match_level=lvl, field=f, field_cn=FIELD_CN[f],
                                      ours=a, edc=b, abs_diff=ad if ad is not None else "", tol=TOL[f],
                                      diff_type=kind, edc_title=title, edc_url=url))
            for io in uo:
                o = ours[io]; st["unmatched_ours"] += 1
                lines.append(dict(base, month=month, check="unmatched_ours", category_ours=o["category"],
                                  tariff_type=o["tariff_type"], voltage_ours=o["voltage"], diff_type="no_edc_row",
                                  ours=o.get("energy_price") if o.get("energy_price") is not None else o.get("flat"),
                                  field="energy_price/flat", edc_title=title, edc_url=url))
            for ie in ue:
                e = erows[ie]; st["unmatched_edc"] += 1
                lines.append(dict(base, month=month, check="unmatched_edc",
                                  category_edc=(e.get("category") or "") + (f" [{e['qual']}]" if e.get("qual") else ""),
                                  tariff_type=e["tariff_type"], voltage_edc=e["voltage"], diff_type="no_our_row",
                                  edc=e.get("energy_price") if e.get("energy_price") is not None else e.get("flat"),
                                  field="energy_price/flat", edc_title=title, edc_url=url,
                                  note=e.get("table_caption", "")[:80]))
        # ---- TOU
        docs = cited_tou_docs(page.get("content"))
        for d in docs:
            edc_docs_seen.setdefault(d, []).append(month)
        edc_rules = parse_tou_text(page.get("content", ""))
        m = int(month[5:])
        er = {"segments": edc_rules[m]} if m in edc_rules else None
        oseg, odoc = our_tou_for_month(con, region, month)
        if er and oseg:
            st["tou_months_compared"] += 1
            eseg = _segs_norm(er["segments"])
            same = eseg == oseg
            no_sharp = _segs_norm([[a, b, "peak" if k == "sharp" else k] for a, b, k in oseg])
            kind = "same" if same else ("same_except_sharp(edc text parse has no 尖峰)"
                                        if no_sharp == eseg and not any(k == "sharp" for _, _, k in eseg) else "different")
            if kind == "different":
                st["tou_months_diff"] += 1
            lines.append(dict(base, month=month, check="tou", diff_type=kind,
                              ours=json.dumps(oseg, ensure_ascii=False), edc=json.dumps(eseg, ensure_ascii=False),
                              edc_title=title, edc_url=url, note="energydc segments = best-effort parse of notes text"))
        elif oseg is None:
            lines.append(dict(base, month=month, check="tou", diff_type="ours_no_tou_rule_for_month",
                              edc=json.dumps(_segs_norm(er["segments"]), ensure_ascii=False) if er else "",
                              edc_title=title, edc_url=url))
    # TOU doc citation line (region level)
    ours_nos = " | ".join(tou_docs_ours)
    edc_nos = " | ".join(f"{d} ({min(ms)}..{max(ms)})" for d, ms in edc_docs_seen.items())
    def _num(d):
        mm = re.search(r"〔\s*(\d{4})\s*〕\s*(\d+)", d)
        return (mm.group(1), mm.group(2)) if mm else None
    on = {_num(d) for d in tou_docs_ours if _num(d)}
    en = {_num(d) for d in edc_docs_seen if _num(d)}
    lines.append(dict(base, check="tou_doc", ours=ours_nos, edc=edc_nos,
                      diff_type=("edc_cites_doc_not_in_ours" if en - on else "consistent_or_subset"),
                      note=("edc-only: " + ", ".join(f"〔{a}〕{b}号" for a, b in sorted(en - on))) if en - on else ""))
    st["tou_doc_ours"] = ours_nos
    st["tou_doc_edc"] = edc_nos
    con.close()
    # write region csv
    os.makedirs(OUT, exist_ok=True)
    cols = list(base.keys())
    with open(os.path.join(OUT, f"{region}.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in lines:
            w.writerow(r)
    st["max_abs_diff"] = round(st["max_abs_diff"], 8)
    json.dump(st, open(os.path.join(OUT, f"{region}.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    log(f"**{region}** (edc `{slug}`): months both={st['months_both']} ours-only={st['months_ours_only']} "
        f"edc-only={st['months_edc_only']}; matched {st['matched_rows']}/{st['ours_rows']} ours rows "
        f"(edc {st['edc_rows']}); rows with value diff {st['rows_with_diff']} (cells {st['value_diffs']}, "
        f"max {st['max_abs_diff']}); null-mismatch cells {st['null_mismatch_cells']}; unmatched ours/edc "
        f"{st['unmatched_ours']}/{st['unmatched_edc']}; TOU months compared {st['tou_months_compared']} "
        f"(diff {st['tou_months_diff']})")
    return st


def rebuild_summary():
    rows = []
    for p in sorted(glob.glob(os.path.join(OUT, "*.json"))):
        st = json.load(open(p, encoding="utf-8"))
        if "region" not in st:
            continue
        rows.append(st)
    cols = ["region", "edc_slug", "months_both", "months_ours_only", "months_edc_only", "matched_rows",
            "rows_with_diff", "max_abs_diff", "value_diffs", "ours_rows", "edc_rows", "unmatched_ours",
            "unmatched_edc", "null_mismatch_cells", "edc_artefact_cells", "tou_months_compared", "tou_months_diff",
            "tou_doc_ours", "tou_doc_edc", "notes", "review"]
    vp = os.path.join(OUT, "review_verdicts.json")
    verdicts = json.load(open(vp, encoding="utf-8")) if os.path.exists(vp) else {}

    def flat(st):
        d = {c: st.get(c, "") for c in cols}
        for c in ("months_ours_only", "months_edc_only"):
            d[c] = f"{len(st[c])}" + (f" ({', '.join(st[c])})" if st[c] else "")
        d["notes"] = " ; ".join(st.get("notes", []))
        d["review"] = verdicts.get(st["region"], "")
        return d
    flat_rows = [flat(s) for s in rows]
    with open(os.path.join(OUT, "summary.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for d in flat_rows:
            w.writerow(d)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active; ws.title = "汇总"
        ws.append(["energydc.cn cross-check – PRIVATE REVIEW ONLY (energydc ToS §4 forbids reuse; not a data source)"])
        ws.append(cols)
        for c in ws[2]:
            c.font = Font(bold=True)
        for d in flat_rows:
            ws.append([d[c] for c in cols])
        ws2 = wb.create_sheet("差异明细")
        ws3 = wb.create_sheet("TOU")
        hdr = None
        red = PatternFill("solid", fgColor="FFF2CC")
        for st in rows:
            p = os.path.join(OUT, f"{st['region']}.csv")
            if not os.path.exists(p):
                continue
            with open(p, encoding="utf-8-sig") as f:
                for r in csv.DictReader(f):
                    if hdr is None:
                        hdr = list(r.keys()); ws2.append(hdr); ws3.append(hdr)
                        for c in ws2[1]: c.font = Font(bold=True)
                        for c in ws3[1]: c.font = Font(bold=True)
                    tgt = ws3 if r["check"] in ("tou", "tou_doc") else ws2
                    if r["check"] == "price" and r["diff_type"] in ("edc_merged_cell_artefact",):
                        continue  # known energydc artefact, kept in region csv only
                    vals = [r[h] for h in hdr]
                    vals = [float(v) if (h in ("ours", "edc", "abs_diff", "tol") and re.fullmatch(r"-?\d+(\.\d+)?(e-?\d+)?", v or "")) else v
                            for h, v in zip(hdr, vals)]
                    tgt.append(vals)
                    if r["diff_type"] == "value":
                        for c in tgt[tgt.max_row]:
                            c.fill = red
        for w_ in (ws, ws2, ws3):
            w_.freeze_panes = "A3" if w_ is ws else "A2"
        wb.save(os.path.join(OUT, "compare_report.xlsx"))
    except Exception as e:
        print("xlsx failed:", e)
    return flat_rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--region", action="append", default=[])
    ap.add_argument("--all-loaded", action="store_true")
    ap.add_argument("--months", default="2025-11..2026-10")
    ap.add_argument("--no-fetch", action="store_true")
    ap.add_argument("--budget", type=float, default=None, help="max seconds of fetching per region")
    ap.add_argument("--rebuild-summary", action="store_true")
    a = ap.parse_args()
    regions = list(a.region)
    if a.all_loaded:
        con = sqlite3.connect(f"file:{pc.DB_PATH}?mode=ro", uri=True, timeout=120)
        regions += [r[0] for r in con.execute("SELECT DISTINCT region_code FROM tariff_month ORDER BY 1")]
        con.close()
    months = month_range(a.months)
    for r in regions:
        try:
            run_region(r, months, fetch=not a.no_fetch, budget=a.budget)
        except Exception as e:
            log(f"**{r}**: ERROR {type(e).__name__}: {str(e)[:300]} – skipped")
    rebuild_summary()


if __name__ == "__main__":
    main()
